from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.config import MAX_FILE_SIZE
from app.errors import AppError
from app.models import DocumentPage, DocumentParagraph, ParsedDocument, SourceType
from app.services.vision import transcribe_image


def _normalize(text: str) -> str:
    return " ".join(text.split())


def _native_blocks(text: str) -> list[DocumentParagraph]:
    return [
        DocumentParagraph(text=normalized, sourceType=SourceType.NATIVE_TEXT)
        for line in text.splitlines()
        if (normalized := _normalize(line))
    ]


def _vision_blocks(payload: dict) -> list[DocumentParagraph]:
    confidence = float(payload["confidence"])
    return [
        DocumentParagraph(text=_normalize(text), sourceType=SourceType.VISION, confidence=confidence)
        for text in payload["paragraphs"]
        if _normalize(text)
    ]


def _merge_unique(blocks: list[DocumentParagraph], additions: list[DocumentParagraph]) -> None:
    existing = {_normalize(block.text) for block in blocks}
    for block in additions:
        normalized = _normalize(block.text)
        if normalized and normalized not in existing:
            blocks.append(block)
            existing.add(normalized)


def _paginate(blocks: list[DocumentParagraph], target_chars: int = 1800) -> list[DocumentPage]:
    pages: list[DocumentPage] = []
    current: list[DocumentParagraph] = []
    chars = 0
    for block in blocks:
        if current and chars + len(block.text) > target_chars:
            pages.append(DocumentPage(page=len(pages) + 1, paragraphs=current))
            current = []
            chars = 0
        current.append(block)
        chars += len(block.text)
    if current:
        pages.append(DocumentPage(page=len(pages) + 1, paragraphs=current))
    return pages


def _docx_text_and_tables(document: Document) -> list[DocumentParagraph]:
    blocks: list[DocumentParagraph] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = _normalize(Paragraph(child, document).text)
            if text:
                blocks.append(DocumentParagraph(text=text, sourceType=SourceType.NATIVE_TEXT))
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            for row in table.rows:
                cells = [_normalize(cell.text) for cell in row.cells]
                text = " | ".join(cell for cell in cells if cell)
                if text:
                    blocks.append(DocumentParagraph(text=text, sourceType=SourceType.TABLE))
    return blocks


async def _parse_docx(filename: str, content: bytes) -> ParsedDocument:
    try:
        document = Document(BytesIO(content))
        blocks = _docx_text_and_tables(document)
        image_parts: list[tuple[bytes, str]] = []
        seen_parts: set[str] = set()
        for relation in document.part.rels.values():
            target = relation.target_part
            content_type = getattr(target, "content_type", "")
            partname = str(getattr(target, "partname", ""))
            if content_type.startswith("image/") and partname not in seen_parts:
                seen_parts.add(partname)
                image_parts.append((target.blob, content_type))
    except Exception as exc:
        raise AppError("parse_failed", "DOCX 文件无法解析，请确认文件未损坏。", 422) from exc
    for image, media_type in image_parts:
        blocks.extend(_vision_blocks(await transcribe_image(image, media_type)))
    if not blocks:
        raise AppError("empty_document", "文档中没有可供审查的文字或图片内容。", 422)
    pages = _paginate(blocks)
    return ParsedDocument(
        name=filename,
        format="DOCX",
        pageCount=len(pages),
        pages=pages,
        text="\n".join(block.text for block in blocks),
        sizeLabel=f"{len(content) / 1024 / 1024:.2f} MB",
    )


async def _parse_pdf(filename: str, content: bytes) -> ParsedDocument:
    try:
        document = fitz.open(stream=content, filetype="pdf")
        if document.needs_pass:
            raise AppError("parse_failed", "PDF 已加密，无法解析。", 422)
    except AppError:
        raise
    except Exception as exc:
        raise AppError("parse_failed", "PDF 文件无法解析，请确认文件未损坏。", 422) from exc
    pages: list[DocumentPage] = []
    try:
        for index, pdf_page in enumerate(document):
            blocks = _native_blocks(pdf_page.get_text("text"))
            image_refs = pdf_page.get_images(full=True)
            native_length = sum(len(block.text) for block in blocks)
            if native_length < 20:
                pixmap = pdf_page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                blocks = _vision_blocks(await transcribe_image(pixmap.tobytes("png"), "image/png"))
            else:
                seen_xrefs: set[int] = set()
                for image_ref in image_refs:
                    xref = image_ref[0]
                    if xref in seen_xrefs:
                        continue
                    seen_xrefs.add(xref)
                    extracted = document.extract_image(xref)
                    media_type = f"image/{extracted.get('ext', 'png')}"
                    _merge_unique(blocks, _vision_blocks(await transcribe_image(extracted["image"], media_type)))
            pages.append(DocumentPage(page=index + 1, paragraphs=blocks))
    finally:
        document.close()
    full_text = "\n".join(block.text for page in pages for block in page.paragraphs)
    if not full_text:
        raise AppError("empty_document", "PDF 中没有识别到可供审查的内容。", 422)
    return ParsedDocument(
        name=filename,
        format="PDF",
        pageCount=len(pages),
        pages=pages,
        text=full_text,
        sizeLabel=f"{len(content) / 1024 / 1024:.2f} MB",
    )


async def parse_document(filename: str, content: bytes) -> ParsedDocument:
    if not content:
        raise AppError("empty_document", "文件为空，请重新选择。", 422)
    if len(content) > MAX_FILE_SIZE:
        raise AppError("file_too_large", "文件超过 20 MB 限制。", 413)
    extension = Path(filename).suffix.lower()
    if extension == ".docx":
        return await _parse_docx(filename, content)
    if extension == ".pdf":
        return await _parse_pdf(filename, content)
    if extension == ".doc":
        raise AppError("unsupported_legacy_word", "暂不支持旧版 DOC，请先转换为 DOCX 后上传。", 415)
    raise AppError("unsupported_format", "仅支持 PDF 和 DOCX 文件。", 415)
