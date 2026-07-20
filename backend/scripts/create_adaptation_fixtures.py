"""Create DOCX fixtures for repeatable parser/upload adaptation checks."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
SOURCE = next(ROOT.joinpath("backend").glob("*冒充客服*.docx"))
OUT = ROOT / "backend" / "test-fixtures"


def _insert_after(paragraph, text: str) -> None:
    new_paragraph = paragraph._parent.add_paragraph(text, style=paragraph.style)
    paragraph._p.addnext(new_paragraph._p)


def _find_paragraph(document: Document, predicate):
    return next((p for p in document.paragraphs if predicate(p.text)), None)


def _set_paragraph(document: Document, index: int, text: str) -> None:
    document.paragraphs[index - 1].text = text


def build() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    # 01: byte-for-byte baseline, preserving the supplied completed document.
    shutil.copy2(SOURCE, OUT / "01-baseline.docx")

    # 02: split a same-paragraph question/answer pair into adjacent paragraphs.
    document = Document(SOURCE)
    paragraph = _find_paragraph(document, lambda text: "问：这是《被害人诉讼权利义务告知书》" in text and "答：" in text)
    if paragraph:
        question, answer = paragraph.text.split("答：", 1)
        paragraph.text = question.rstrip()
        _insert_after(paragraph, "答：" + answer.strip())
    document.save(OUT / "02-line-breaks.docx")

    # 03: retain an explicit answer marker but remove selected answer content.
    document = Document(SOURCE)
    paragraph = _find_paragraph(document, lambda text: text.strip().startswith("答：") and "反诈宣传" in text)
    if paragraph:
        paragraph.text = "答："
    document.save(OUT / "03-blank-answer.docx")

    # 04: force logical pagination and long-text normalization.
    document = Document(SOURCE)
    paragraph = _find_paragraph(document, lambda text: "分两笔通过手机银行转账共计" in text)
    if paragraph:
        paragraph.text += (
            " 其后我又反复核对了聊天记录、来电记录、手机银行流水和设备通知。"
            "这段补充说明用于验证长段落、中文标点、数字金额、引号、括号以及跨逻辑页的证据定位。" * 35
        )
    document.save(OUT / "04-long-answer.docx")

    # 05: add structured table blocks plus punctuation and mixed scalar values.
    document = Document(SOURCE)
    document.add_paragraph("补充材料：以下表格为脱敏测试数据，不替代正文问答。")
    table = document.add_table(rows=1, cols=3)
    # The supplied template does not define the optional "Table Grid" style.
    # Keep the table unstyled so this fixture tests structure, not style names.
    for cell, value in zip(table.rows[0].cells, ("时间", "渠道", "金额/备注")):
        cell.text = value
    for row in (
        ("2026-07-16 16:40", "企业聊天账号", "0 元（引流）"),
        ("2026-07-16 16:58", "手机银行", "18,200.00 元"),
        ("2026-07-16 17:05", "手机银行", "18,200.00 元；合计 36,400 元"),
        ("2026-07-16 17:12", "报警", "备注：客服退款 / 远程协助 APP"),
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    document.save(OUT / "05-table-and-symbols.docx")

    # 06: demonstration record with all review statuses represented.
    document = Document(SOURCE)
    _set_paragraph(document, 26, "答：")
    _set_paragraph(document, 18, "答：")
    _set_paragraph(
        document,
        32,
        "答：2026年7月16日13时42分，我在测试省测试市新城区平安街道清风社区家中接到冒充客服来电。"
        "我通过本人实名手机号与对方普通电话联系3次。对方发送测试链接让我下载远程协助APP并开启屏幕共享。14时26分和14时51分，"
        "我分别通过手机银行转账12800元、23600元，线上转账共36400元。"
        "15时10分，对方又预约我到测试银行新城支行取现5000元，并在该支行门口将现金交给自称理赔员的人。"
        "我未收到任何返款，15时32分挂失账户并报警。",
    )
    _set_paragraph(
        document,
        44,
        "答：",
    )
    _set_paragraph(
        document,
        46,
        "答：对方预约我于2026年7月16日15时10分，到测试银行新城支行门口交付5000元现金。",
    )
    _set_paragraph(
        document,
        48,
        "答：取款前我通过测试银行客服预约取现5000元；银行于15时02分电话提示防范诈骗风险。",
    )
    _set_paragraph(
        document,
        50,
        "答：我于2026年7月16日15时06分在测试银行新城支行（测试省测试市新城区平安路88号）柜台取现5000元，共1次。",
    )
    _set_paragraph(document, 87, "其它形式：屏幕共享。")
    _set_paragraph(
        document,
        40,
        "答：我使用本人实名登记的测试手机号15500000018，在13时42分、13时50分和14时02分通过普通电话与对方联系3次；"
        "14时02分收到运营商关于陌生客服来电的风险提示短信。",
    )
    _set_paragraph(document, 42, "答：仅留有一张模糊截图，无法判断我和对方是否使用过聊天软件，也无法核验账号是否实名。")
    _set_paragraph(document, 70, "其它：无。")
    _set_paragraph(document, 71, "无可核验的聊天账号、昵称或聊天记录。")
    _set_paragraph(
        document,
        100,
        "答：被骗总额36400元，收到返款0元，但我认为实际净损失为30000元。",
    )
    _set_paragraph(
        document,
        127,
        "答：☑ 现金。2026年7月16日15时06分，我从本人测试银行账户取现5000元。",
    )
    _set_paragraph(document, 128, "资金来源：本人测试银行借记卡账户余额；柜台取款1次，金额5000元。")
    _set_paragraph(document, 129, "与嫌疑人接触共1次。")
    _set_paragraph(document, 130, "第1次交接：向自称“理赔员”的人员交付5000元现金。")
    _set_paragraph(document, 131, "时间：2026年7月16日15时10分。")
    _set_paragraph(document, 132, "地点：测试银行新城支行门口（测试省测试市新城区平安路88号）。")
    _set_paragraph(document, 133, "金额：5000元。")
    _set_paragraph(document, 135, "☐ 线下邮寄；本次未使用邮寄或物流。")
    _set_paragraph(document, 136, "☑ 面对面交付；接收人为自称“理赔员”的人员，未留下真实身份信息。")
    _set_paragraph(document, 137, "其它：☐ 同城配送  ☐ 上门取现  ☐ 其他。")
    _set_paragraph(document, 139, "答：")
    document.save(OUT / "06-all-statuses-demo.docx")


if __name__ == "__main__":
    build()
    print(f"created 6 fixtures in {OUT}")
