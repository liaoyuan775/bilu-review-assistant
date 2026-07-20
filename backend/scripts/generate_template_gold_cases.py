from __future__ import annotations

import argparse
from html import escape
import json
from pathlib import Path
import re
import sys
from typing import Iterable
from docx import Document
import fitz
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.data.rules import TEMPLATE_RULE_CATALOG  # noqa: E402
from app.review.extraction import DOMAIN_ENTITY_FIELDS, DOMAIN_FACT_PATHS  # noqa: E402
from app.review.rules import evaluate_template_rules  # noqa: E402
from app.core.template_models import CaseExtraction, ExtractedEntity, ExtractedFact  # noqa: E402


PROHIBITED_SENSITIVE_PATTERNS = {
    "identity_number": re.compile(r"(?<!\d)\d{17}[0-9Xx](?!\d)"),
    "phone_number": re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"),
    "bank_card": re.compile(r"(?<!\d)(?:62\d{14,17}|[3-6]\d{15})(?!\d)"),
    "url": re.compile(r"https?://", re.IGNORECASE),
    "public_ip": re.compile(
        r"(?<!\d)(?:(?:[1-9]|[1-9]\d|1\d\d|2[0-4]\d|25[0-5])\.){3}"
        r"(?:[1-9]|[1-9]\d|1\d\d|2[0-4]\d|25[0-5])(?!\d)"
    ),
}


SCENARIOS = (
    ("gold-01-procedure", "权利义务和笔录核对", ("header_procedure",)),
    ("gold-02-timeline", "接触转账报案时间线", ("case_timeline",)),
    ("gold-03-contact-switch", "多平台切换", ("contact_channels",)),
    ("gold-04-risk-warning", "支付风险提示与证据", ("risk_and_evidence",)),
    ("gold-05-online-money", "线上转账与返利", ("online_money",)),
    ("gold-06-offline-delivery", "线下现金与财物交付", ("offline_delivery",)),
    ("gold-07-extra", "其他补充事实", ("case_timeline",)),
    ("gold-08-case-procedure", "身份程序与案情概述", ("header_procedure", "case_timeline")),
    ("gold-09-contact-risk", "联系渠道与风险提示", ("contact_channels", "risk_and_evidence")),
    ("gold-10-online-timeline", "线上资金与时间线", ("case_timeline", "online_money")),
    ("gold-11-offline-timeline", "线下交付与时间线", ("case_timeline", "offline_delivery")),
    ("gold-12-complete", "全业务域完整问询", tuple(DOMAIN_FACT_PATHS)),
)


NATURAL_CASES = (
    {
        "id": "natural-01-complete-online",
        "name": "冒充客服两笔转账完整询问",
        "mutationKind": "complete",
        "paragraphs": (
            "问：你为什么来报案，事情发生在哪里？",
            "答：我在星桥区青禾街道松林社区家中遇到冒充客服退款诈骗，实际损失二百八十元。",
            "问：对方最早何时、通过什么方式联系你？",
            "答：二〇二六年七月一日上午九点，对方先发短信，说我购买的商品有质量问题，可以办理退款。",
            "问：你把整个付款经过按顺序讲清楚。",
            "答：当天上午十点我按对方要求转了一百元，十一点又转了二百元，后来收到二十元返款，两笔共转出三百元，净损失二百八十元。",
            "问：两笔钱分别转到哪里，用什么方式支付？",
            "答：第一笔上午十点从测试付款账户甲通过手机银行转到测试收款账户甲，金额一百元，交易流水为测试流水甲；第二笔上午十一点从测试付款账户乙扫对方发来的收款码支付到测试收款账户乙，金额二百元，交易流水为测试流水乙。",
            "问：转账是你本人操作的吗，对方有没有远程控制设备，你是否透露过验证码？",
            "答：两笔都是我本人操作，对方没有远程控制，但我把短信验证码告诉了对方。",
        ),
        "expectedRuleStatuses": {
            "MONEY-001": "covered",
            "MONEY-002": "covered",
            "MONEY-003": "covered",
            "MONEY-004": "covered",
        },
        "expectedFacts": {
            "money.gross_loss": {"value": 300, "clarity": "clear"},
            "money.rebate_total": {"value": 20, "clarity": "clear"},
            "money.net_loss": {"value": 280, "clarity": "clear"},
            "money.remote_control_used": {"value": False, "clarity": "clear"},
            "money.credentials_disclosed": {"value": True, "clarity": "clear"},
            "online_money.transfer_count": {"value": 2, "clarity": "clear"},
            "online_money.total": {"value": 300, "clarity": "clear"},
        },
        "expectedEntities": {
            "transfers": [
                {"amount": {"value": 100, "clarity": "clear"}},
                {"amount": {"value": 200, "clarity": "clear"}},
            ],
        },
    },
    {
        "id": "natural-02-missing-contact",
        "name": "未询问首次接触方式",
        "mutationKind": "missing",
        "paragraphs": (
            "问：你是否听清楚应当如实陈述？",
            "答：听清楚了，我会如实陈述。",
            "问：你还有其他需要补充的吗？",
            "答：暂时没有。",
        ),
        "expectedRuleStatuses": {"LEAD-001": "missing"},
        "expectedFacts": {},
        "expectedEntities": {},
    },
    {
        "id": "natural-03-unclear-warning",
        "name": "支付风险提示回答不清",
        "mutationKind": "unclear",
        "paragraphs": (
            "问：你一共向对方转过几次钱？",
            "答：转过两次。",
            "问：付款时是否收到过风险提示，是通过什么渠道提示了什么内容？",
            "答：我记不清有没有收到，也说不清是什么渠道和什么内容。",
        ),
        "expectedRuleStatuses": {"RISK-001": "incomplete"},
        "expectedFacts": {
            "risk.payment_warning_received": {"value": None, "clarity": "unclear"},
            "risk.payment_warning_channel": {"value": None, "clarity": "unclear"},
            "risk.payment_warning_content": {"value": None, "clarity": "unclear"},
        },
        "expectedEntities": {},
    },
    {
        "id": "natural-04-inconsistent-total",
        "name": "转账明细与总额矛盾",
        "mutationKind": "inconsistent",
        "paragraphs": (
            "问：你线上一共转了几笔、总共多少钱？",
            "答：一共两笔，总共三百元。",
            "问：把每一笔的金额和去向分别讲清楚。",
            "答：第一笔一百元转到测试账户甲，第二笔一百五十元转到测试账户乙。",
        ),
        "expectedRuleStatuses": {"MONEY-004": "inconsistent"},
        "expectedFacts": {
            "online_money.transfer_count": {"value": 2, "clarity": "clear"},
            "online_money.total": {"value": 300, "clarity": "clear"},
        },
        "expectedEntities": {
            "transfers": [
                {"amount": {"value": 100, "clarity": "clear"}},
                {"amount": {"value": 150, "clarity": "clear"}},
            ],
        },
    },
)


def _value_for(path: str):
    if path in {
        "money.gross_loss", "online_money.total",
    }:
        return 300
    if path == "money.rebate_total":
        return 20
    if path == "money.net_loss":
        return 280
    if path in {
        "online_money.transfer_count", "contact.switch_count",
    }:
        return 2
    if path in {"cash.withdrawal_count", "offline.handoff_count"}:
        return 1
    if path == "timeline.first_contact_at":
        return "2026-07-01T09:00:00"
    if path in {"timeline.incident_at", "privacy.disclosure_time", "risk.loss_report_time"}:
        return "2026-07-01T10:00:00"
    boolean_markers = (
        ".used", ".received", ".installed", ".exists", ".available",
        ".occurred", ".requested", ".reviewed", ".confirmed", ".reported",
        ".can_provide", ".chat_used", ".phone_used", ".voice_call",
    )
    if path.endswith(boolean_markers) or path in {
        "procedure.key_information_reconfirmed", "procedure.record_matches_statement",
        "procedure.record_reviewed", "procedure.rights_notice_read", "procedure.statement_confirmed_true",
        "procedure.truth_notice_confirmed", "prevention.community_police_publicity",
        "risk.phone_card_real_name", "risk.suspect_chat_real_name",
        "risk.victim_chat_real_name", "money.credentials_disclosed",
        "money.remote_control_used", "offline.suspect_appointment",
    }:
        return True
    return f"VALUE_TEST_{path.upper().replace('.', '_')}"


def _question_answers(domains: tuple[str, ...]) -> list[tuple[str, str]]:
    def display(value) -> str:
        if value is True:
            return "是，已明确发生或确认"
        if value is False:
            return "否，已明确未发生"
        return str(value)

    pairs = [
        (
            f"请明确记录字段 {path}。",
            f"字段 {path} 的脱敏测试值为 {display(_value_for(path))}，该回答明确。",
        )
        for domain in domains
        for path in DOMAIN_FACT_PATHS[domain]
    ]
    for entity_type, fields in {
        name: fields for domain in domains for name, fields in DOMAIN_ENTITY_FIELDS[domain].items()
    }.items():
        count = 2 if entity_type in {"transfers", "contact_switches"} else 1
        for index in range(1, count + 1):
            values = []
            for field in fields:
                if entity_type == "transfers" and field == "amount":
                    value = index * 100
                elif entity_type == "rebates" and field == "amount":
                    value = 20
                elif entity_type == "withdrawals" and field == "amount":
                    value = 300
                elif entity_type == "offline_handoffs" and field == "amount_or_value":
                    value = 300
                elif field == "time":
                    value = f"2026-07-01T1{index}:00:00"
                else:
                    value = f"ENTITY_TEST_{entity_type.upper()}_{index}_{field.upper()}"
                values.append(f"{field}={value}")
            pairs.append((
                f"请逐项记录 {entity_type} 第 {index} 项。",
                f"{entity_type} 第 {index} 项明确记录为：" + "；".join(values) + "。",
            ))
    return pairs


def _case_extraction(domains: tuple[str, ...]) -> CaseExtraction:
    facts = {
        path: ExtractedFact(
            value=_value_for(path), clarity="clear",
            evidenceAnchorIds=[f"gold-{path}"], sourceConfidence=1.0,
        )
        for domain in domains
        for path in DOMAIN_FACT_PATHS[domain]
    }
    if "online_money" in domains and "case_timeline" not in domains:
        facts["timeline.incident_at"] = ExtractedFact(
            value="2026-07-01T11:00:00", clarity="clear",
            evidenceAnchorIds=["gold-transfers-1-time"], sourceConfidence=1.0,
        )
    entities: dict[str, list[ExtractedEntity]] = {}
    for domain in domains:
        for entity_type, fields in DOMAIN_ENTITY_FIELDS[domain].items():
            count = 2 if entity_type in {"transfers", "contact_switches"} else 1
            items = []
            for index in range(1, count + 1):
                values = {}
                for field in fields:
                    if entity_type == "transfers" and field == "amount":
                        value = index * 100
                    elif entity_type == "rebates" and field == "amount":
                        value = 20
                    elif entity_type == "withdrawals" and field == "amount":
                        value = 300
                    elif entity_type == "offline_handoffs" and field == "amount_or_value":
                        value = 300
                    elif field == "time":
                        value = f"2026-07-01T1{index}:00:00"
                    else:
                        value = f"ENTITY_TEST_{entity_type.upper()}_{index}_{field.upper()}"
                    values[field] = ExtractedFact(
                        value=value, clarity="clear",
                        evidenceAnchorIds=[f"gold-{entity_type}-{index}-{field}"],
                        sourceConfidence=1.0,
                    )
                items.append(ExtractedEntity(
                    id=f"{entity_type}-{index}", entityType=entity_type, fields=values,
                ))
            entities[entity_type] = items
    return CaseExtraction(facts=facts, entities=entities)


def _write_docx(
    path: Path,
    scenario_name: str,
    questions: list[tuple[str, str]],
    domains: tuple[str, ...],
) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = 210 * 36000, 297 * 36000
    document.add_heading("电信网络诈骗询问笔录金标准", 0)
    document.add_paragraph(f"虚构脱敏测试材料 | {scenario_name} | 禁止用于真实案件")
    if "offline_delivery" in domains:
        document.add_paragraph(
            "本案明确发生线下取现和财物交付，offline.used 为是；仅有 1 次实际取款和 1 次实际交付，预约信息不是另一笔取款。"
        )
    if "online_money" in domains:
        document.add_paragraph(
            "本案明确发生线上转账和返款，online_money.used 为是；共有 2 笔转账、1 笔返款。"
        )
    if "header_procedure" in domains:
        document.add_paragraph("案件编号：CASE_TEST_GOLD；被询问人：PERSON_TEST；联系电话：PHONE_TEST；身份证：ID_TEST；银行卡：CARD_TEST。")
    for index, (question, answer) in enumerate(questions, start=1):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{index:03d} 问：{question}").bold = True
        paragraph.add_run().add_break()
        paragraph.add_run(f"答：{answer}")
    if "header_procedure" in domains:
        document.add_paragraph("被询问人已核对以上笔录，确认记录与陈述一致。签名：SIGNATURE_TEST。")
    document.save(path)


def _font_name() -> str:
    name = "GoldCorpusCJK"
    if name not in pdfmetrics.getRegisteredFontNames():
        font_path = Path("C:/Windows/Fonts/msyh.ttc")
        if font_path.exists():
            pdfmetrics.registerFont(TTFont(name, str(font_path), subfontIndex=0))
        else:
            name = "STSong-Light"
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(UnicodeCIDFont(name))
    return name


def _write_pdf(
    path: Path,
    scenario_name: str,
    questions: list[tuple[str, str]],
    domains: tuple[str, ...],
) -> None:
    font = _font_name()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("GoldTitle", parent=styles["Title"], fontName=font, fontSize=18, leading=24, alignment=TA_CENTER)
    body = ParagraphStyle("GoldBody", parent=styles["BodyText"], fontName=font, fontSize=9, leading=14, spaceAfter=4)
    story = [
        Paragraph("电信网络诈骗询问笔录金标准", title),
        Paragraph(escape(f"虚构脱敏测试材料 | {scenario_name} | 禁止用于真实案件"), body),
        Spacer(1, 4 * mm),
    ]
    if "offline_delivery" in domains:
        story.append(Paragraph(
            "本案明确发生线下取现和财物交付，offline.used 为是；仅有 1 次实际取款和 1 次实际交付，预约信息不是另一笔取款。",
            body,
        ))
    if "online_money" in domains:
        story.append(Paragraph(
            "本案明确发生线上转账和返款，online_money.used 为是；共有 2 笔转账、1 笔返款。",
            body,
        ))
    for index, (question, answer) in enumerate(questions, start=1):
        story.append(Paragraph(
            f"{index:03d} 问：{escape(question)}<br/>答：{escape(answer)}",
            body,
        ))
    if "header_procedure" in domains:
        story.extend([PageBreak(), Paragraph("被询问人已核对笔录并确认记录与陈述一致。签名：SIGNATURE_TEST。", body)])
    SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
    ).build(story)


def _write_natural_docx(path: Path, case: dict) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = 210 * 36000, 297 * 36000
    document.add_heading("电信网络诈骗询问笔录脱敏测试件", 0)
    document.add_paragraph(f"虚构材料 | {case['name']} | 禁止用于真实案件")
    for paragraph in case["paragraphs"]:
        document.add_paragraph(paragraph)
    document.save(path)


def _write_natural_pdf(path: Path, case: dict) -> None:
    font = _font_name()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("NaturalTitle", parent=styles["Title"], fontName=font, fontSize=18, leading=24, alignment=TA_CENTER)
    body = ParagraphStyle("NaturalBody", parent=styles["BodyText"], fontName=font, fontSize=10, leading=16, spaceAfter=5)
    story = [
        Paragraph("电信网络诈骗询问笔录脱敏测试件", title),
        Paragraph(escape(f"虚构材料 | {case['name']} | 禁止用于真实案件"), body),
        Spacer(1, 4 * mm),
        *(Paragraph(escape(paragraph), body) for paragraph in case["paragraphs"]),
    ]
    SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=16 * mm,
    ).build(story)


def _oracle() -> dict:
    cases = []
    for case_id, name, domains in SCENARIOS:
        extraction = _case_extraction(domains)
        issues = evaluate_template_rules(extraction, TEMPLATE_RULE_CATALOG.rules)
        cases.append({
            "id": case_id,
            "name": name,
            "docx": f"docx/{case_id}.docx",
            "pdf": f"pdf/{case_id}.pdf",
            "domains": list(domains),
            "questionCount": len(_question_answers(domains)),
            "expectedRuleStatuses": {issue.ruleId: issue.status.value for issue in issues},
        })
    mutations = []
    for rule in TEMPLATE_RULE_CATALOG.rules:
        for kind in ("covered", "missing", "incomplete", "inconsistent"):
            expected = kind
            if kind == "inconsistent" and not rule.consistencyChecks:
                expected = "incomplete"
            if (
                kind in {"missing", "incomplete"}
                and rule.appliesWhen is not None
                and rule.appliesWhen.path in rule.requiredFields
            ):
                expected = "needs_manual_review"
            mutations.append({
                "ruleId": rule.ruleId,
                "kind": kind,
                "expectedStatus": expected,
            })
    natural_cases = [{
        "id": case["id"],
        "name": case["name"],
        "mutationKind": case["mutationKind"],
        "docx": f"natural/docx/{case['id']}.docx",
        "pdf": f"natural/pdf/{case['id']}.pdf",
        "questionCount": sum(paragraph.startswith("问：") for paragraph in case["paragraphs"]),
        "expectedRuleStatuses": case["expectedRuleStatuses"],
        "expectedFacts": case["expectedFacts"],
        "expectedEntities": case["expectedEntities"],
    } for case in NATURAL_CASES]
    return {
        "schemaVersion": 1,
        "ruleCatalogVersion": TEMPLATE_RULE_CATALOG.version,
        "naturalOracleSource": "hand-authored-police-review-v1",
        "naturalCases": natural_cases,
        "cases": cases,
        "mutations": mutations,
    }


def build_gold_corpus(output_dir: Path) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "docx").mkdir(exist_ok=True)
    (output_dir / "pdf").mkdir(exist_ok=True)
    (output_dir / "natural" / "docx").mkdir(parents=True, exist_ok=True)
    (output_dir / "natural" / "pdf").mkdir(parents=True, exist_ok=True)
    oracle = _oracle()
    oracle_json = json.dumps(oracle, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    for case in oracle["cases"]:
        domains = tuple(case["domains"])
        questions = _question_answers(domains)
        _write_docx(output_dir / case["docx"], case["name"], questions, domains)
        _write_pdf(output_dir / case["pdf"], case["name"], questions, domains)
    natural_source = {case["id"]: case for case in NATURAL_CASES}
    for case in oracle["naturalCases"]:
        source = natural_source[case["id"]]
        _write_natural_docx(output_dir / case["docx"], source)
        _write_natural_pdf(output_dir / case["pdf"], source)
    (output_dir / "template_cases.json").write_text(
        oracle_json, encoding="utf-8", newline="\n"
    )
    return {**oracle, "oracleJson": oracle_json}


def _extract_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        for section in document.sections:
            parts.extend(paragraph.text for paragraph in section.header.paragraphs)
            parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
        return "\n".join(parts)
    if path.suffix.lower() == ".pdf":
        with fitz.open(path) as document:
            return "\n".join(page.get_text() for page in document)
    return path.read_text(encoding="utf-8", errors="ignore")


def scan_prohibited_sensitive_shapes(paths: Iterable[Path]) -> list[dict[str, str]]:
    findings: list[dict[str, str]] = []
    for path in paths:
        text = _extract_text(path)
        for name, pattern in PROHIBITED_SENSITIVE_PATTERNS.items():
            match = pattern.search(text)
            if match:
                findings.append({"file": path.name, "pattern": name})
    return findings


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate sanitized template-review gold records.")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--oracle", type=Path)
    args = parser.parse_args()
    corpus = build_gold_corpus(args.output_dir)
    if args.oracle:
        args.oracle.parent.mkdir(parents=True, exist_ok=True)
        args.oracle.write_text(corpus["oracleJson"], encoding="utf-8", newline="\n")
    print(json.dumps({
        "contractCases": len(corpus["cases"]),
        "naturalCases": len(corpus["naturalCases"]),
        "mutations": len(corpus["mutations"]),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
