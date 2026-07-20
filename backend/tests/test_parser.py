import asyncio
import base64
from io import BytesIO
from unittest.mock import AsyncMock

from docx import Document

from app.parsing.parser import _native_blocks, _parse_docx


def _docx_with_image_between_question_answers() -> bytes:
    image = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    document = Document()
    document.add_paragraph("问：是否保留转账凭证？")
    document.add_paragraph("答：凭证见下图。")
    document.add_picture(BytesIO(image))
    document.add_paragraph("问：是否核对笔录？")
    document.add_paragraph("答：已经核对。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_native_pdf_lines_merge_into_complete_question_answer_pairs():
    blocks = _native_blocks(
        """页眉
03
问：请按时间顺序讲述事情经过。
答：2026年7月15日19时32分，我在家中看到广告，随后添加对
方微信。对方让我下载APP并转账，我在20时18分至21时06分之
间共转账三笔，之后意识到被骗并报警。
04
问：对方通过哪些平台联系你？
答：对方先通过短视频平台联系，后来添加微信。"""
    )

    assert [block.text for block in blocks] == [
        "页眉",
        "03 问：请按时间顺序讲述事情经过。 "
        "答：2026年7月15日19时32分，我在家中看到广告，随后添加对"
        "方微信。对方让我下载APP并转账，我在20时18分至21时06分之"
        "间共转账三笔，之后意识到被骗并报警。",
        "04 问：对方通过哪些平台联系你？ 答：对方先通过短视频平台联系，后来添加微信。",
    ]


def test_native_pdf_question_pair_stops_before_closing_and_signatures():
    blocks = _native_blocks(
        """15
问：以上内容是否确认？
答：我确认以上内容无误。
询问结束后，被询问人逐页核对笔录。
被询问人签名：张测试
核对结果：与陈述一致
本文件全部信息均为虚构。"""
    )

    assert [block.text for block in blocks] == [
        "15 问：以上内容是否确认？ 答：我确认以上内容无误。",
        "询问结束后，被询问人逐页核对笔录。",
        "被询问人签名：张测试",
        "核对结果：与陈述一致",
        "本文件全部信息均为虚构。",
    ]


def test_docx_image_transcription_stays_with_the_preceding_question_answer(monkeypatch):
    transcribe = AsyncMock(return_value={
        "paragraphs": ["图片内容：转账凭证流水号 TEST-INLINE-001。"],
        "confidence": 0.95,
    })
    monkeypatch.setattr("app.parsing.parser.transcribe_image", transcribe)

    parsed = asyncio.run(_parse_docx(
        "qa-with-inline-image.docx",
        _docx_with_image_between_question_answers(),
    ))

    assert transcribe.await_count == 1
    assert len(parsed.questionAnswers) == 2
    assert "TEST-INLINE-001" in parsed.questionAnswers[0].answer
    assert "TEST-INLINE-001" not in parsed.questionAnswers[1].answer
