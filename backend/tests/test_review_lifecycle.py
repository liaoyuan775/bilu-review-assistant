from copy import deepcopy
import asyncio
from io import BytesIO
from unittest.mock import AsyncMock

import pytest
from docx import Document

from app.core.errors import AppError
from app.core.models import (
    ArtifactSummary,
    DocumentPage,
    DocumentParagraph,
    DocumentWarning,
    ManualDecision,
    ManualStatus,
    ParsedDocument,
    ReviewMode,
    ReviewResult,
    ReviewStatus,
    ReviewTask,
    RuleStatus,
    SourceType,
    TaskStatus,
)
from app.reporting.archive import assert_archive_ready
from app.review import review, extraction
from app.review.extraction import (
    DOMAIN_ENTITY_FIELDS,
    DOMAIN_FACT_PATHS,
    DOMAIN_ORDER,
    TemplateReviewOutcome,
)
from app.storage.artifacts import ArtifactStorage
from app.storage import artifacts
from app.storage import store
from app.storage.store import SqliteTaskStore
from app.core.template_models import CaseExtraction, ExtractedFact, TemplateReviewIssue
from app.main import app


REQUIRED_ARTIFACTS = ["review_pdf", "follow_up_docx", "structured_json", "archive_manifest"]


def _task() -> ReviewTask:
    document = ParsedDocument(
        name="脱敏测试笔录.docx",
        format="DOCX",
        pageCount=1,
        pages=[DocumentPage(page=1, paragraphs=[
            DocumentParagraph(id="b1", text="问：测试？答：明确。", sourceType=SourceType.NATIVE_TEXT),
        ])],
        text="问：测试？答：明确。",
        sizeLabel="test",
    )
    result = ReviewResult(
        ruleId="CASE-001",
        ruleName="测试规则",
        category="CASE",
        group="CASE",
        status=RuleStatus.COVERED,
        missingFacts=[],
        evidence="明确证据",
        evidenceLocation=None,
        reason="已覆盖。",
        suggestedQuestion="",
        source="内部模板测试",
        severity="high",
    )
    return ReviewTask(
        mode=ReviewMode.QWEN,
        status=TaskStatus.COMPLETED,
        document=document,
        documentId="document-1",
        documentVersionId="version-1",
        reviewRunId="run-1",
        results=[result],
        requiredArtifacts=REQUIRED_ARTIFACTS,
        artifacts=[
            ArtifactSummary(id=f"artifact-{kind}", type=kind, filename=f"{kind}.bin", sha256="a" * 64, sizeBytes=10)
            for kind in REQUIRED_ARTIFACTS
        ],
    )


def _complete_extraction() -> CaseExtraction:
    return CaseExtraction(
        facts={
            path: ExtractedFact(value=None, clarity="missing", evidenceAnchorIds=[])
            for domain in DOMAIN_ORDER
            for path in DOMAIN_FACT_PATHS[domain]
        },
        entities={
            entity_type: []
            for domain in DOMAIN_ORDER
            for entity_type in DOMAIN_ENTITY_FIELDS[domain]
        },
    )


@pytest.mark.parametrize(
    ("mutate", "expected_code"),
    [
        (lambda task: setattr(task, "status", TaskStatus.FAILED), "archive_task_incomplete"),
        (lambda task: task.failedDomains.append("case_timeline"), "archive_failed_domains"),
        (
            lambda task: task.document.warnings.append(DocumentWarning(code="media_corrupt", message="测试告警")),
            "archive_unresolved_warnings",
        ),
        (lambda task: task.artifacts.pop(), "archive_missing_artifacts"),
        (
            lambda task: task.results.__setitem__(0, task.results[0].model_copy(update={
                "status": RuleStatus.MISSING,
                "manualDecision": ManualDecision(status=ManualStatus.PENDING),
            })),
            "archive_pending_high_risk",
        ),
    ],
)
def test_archive_gates_have_precise_error_codes(mutate, expected_code):
    task = _task()
    mutate(task)

    with pytest.raises(AppError) as error:
        assert_archive_ready(task)

    assert error.value.code == expected_code


def test_acknowledged_warning_and_resolved_high_risk_issue_can_archive():
    task = _task()
    task.document.warnings.append(DocumentWarning(code="media_corrupt", message="测试告警"))
    task.acknowledgedWarnings.append("media_corrupt")
    task.results[0] = task.results[0].model_copy(update={
        "status": RuleStatus.INCOMPLETE,
        "manualDecision": ManualDecision(status=ManualStatus.RESOLVED, reason="已补问并记录明确答案。"),
    })

    assert_archive_ready(deepcopy(task))


def test_issue_action_appends_event_and_archived_task_is_read_only(tmp_path, monkeypatch):
    database = tmp_path / "reviews.db"
    monkeypatch.setattr(store, "STORE", SqliteTaskStore(database))
    task = _task()
    task.results[0] = task.results[0].model_copy(update={
        "status": RuleStatus.MISSING,
        "manualDecision": ManualDecision(status=ManualStatus.PENDING),
    })
    store.save_task(task)

    result = review.record_issue_action(
        task.id,
        "CASE-001",
        ManualStatus.RESOLVED,
        "已核对补问答案。",
        "test-operator",
    )

    assert result.manualDecision.status == ManualStatus.RESOLVED
    events = store.STORE.get_audit_snapshot(task.id)["events"]
    assert [(event["event_type"], event["actor_id"]) for event in events] == [
        ("resolved", "test-operator"),
    ]

    task = store.get_task(task.id)
    task.reviewStatus = ReviewStatus.ARCHIVED
    store.save_task(task)
    with pytest.raises(AppError) as error:
        review.record_issue_action(
            task.id,
            "CASE-001",
            ManualStatus.IGNORED,
            "不能修改。",
            "test-operator",
        )
    assert error.value.code == "review_archived"


def test_follow_up_answer_creates_version_event_and_affected_domain_rerun(tmp_path, monkeypatch):
    database = tmp_path / "reviews.db"
    monkeypatch.setattr(store, "STORE", SqliteTaskStore(database))
    task = _task()
    task.results[0] = task.results[0].model_copy(update={
        "status": RuleStatus.MISSING,
        "manualDecision": ManualDecision(status=ManualStatus.SUPPLEMENTED, reason="请补充报案原因。"),
    })
    task.extractionPayload = CaseExtraction().model_dump(mode="json")
    store.save_task(task)
    task.documentId = store.STORE.save_document(
        task.id,
        filename=task.document.name,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        sha256="a" * 64,
        original_path="test/original.docx",
    )
    task.documentVersionId = store.STORE.save_document_version(
        task.documentId,
        content_sha256="a" * 64,
        parsed_payload=task.document.model_dump(mode="json"),
    )
    store.save_task(task)
    resolved_result = task.results[0].model_copy(update={"status": RuleStatus.COVERED})
    outcome = TemplateReviewOutcome(
        extraction=CaseExtraction(),
        issues=[],
        results=[resolved_result],
    )
    run = AsyncMock(return_value=outcome)
    monkeypatch.setattr(review, "run_template_review", run)

    updated = asyncio.run(review.record_follow_up_answer(
        task.id,
        "CASE-001",
        question="你因何事报案？",
        answer="因测试诈骗损失报案。",
        actor_id="test-operator",
    ))

    assert updated.documentVersionId != task.documentVersionId
    assert updated.results[0].manualDecision.status == ManualStatus.RESOLVED
    assert run.await_args.kwargs["domains"] == ("case_timeline",)
    snapshot = store.STORE.get_audit_snapshot(task.id)
    assert len(snapshot["versions"]) == 2
    assert snapshot["events"][-1]["event_type"] == "follow_up_answer"


def test_follow_up_preserves_manual_decisions_from_unaffected_domains(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "STORE", SqliteTaskStore(tmp_path / "reviews.db"))
    task = _task()
    task.results.extend([
        task.results[0].model_copy(update={
            "ruleId": "CASE-002",
            "manualDecision": ManualDecision(
                status=ManualStatus.CONFIRMED,
                reason="原时间线问题已确认。",
            ),
        }),
        task.results[0].model_copy(update={
            "ruleId": "CONTACT-001",
            "group": "CONTACT",
            "category": "CONTACT",
            "manualDecision": ManualDecision(
                status=ManualStatus.IGNORED,
                reason="已人工核对联系方式。",
            ),
        }),
    ])
    task.extractionPayload = _complete_extraction().model_dump(mode="json")
    store.save_task(task)
    task.documentId = store.STORE.save_document(
        task.id,
        filename=task.document.name,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        sha256="a" * 64,
        original_path="test/original.docx",
    )
    task.documentVersionId = store.STORE.save_document_version(
        task.documentId,
        content_sha256="a" * 64,
        parsed_payload=task.document.model_dump(mode="json"),
    )
    store.save_task(task)
    refreshed = [
        item.model_copy(update={"manualDecision": ManualDecision()})
        for item in task.results
    ]
    monkeypatch.setattr(review, "run_template_review", AsyncMock(return_value=TemplateReviewOutcome(
        extraction=_complete_extraction(),
        issues=[],
        results=refreshed,
    )))

    updated = asyncio.run(review.record_follow_up_answer(
        task.id,
        "CASE-001",
        question="你因何事报案？",
        answer="因测试诈骗损失报案。",
        actor_id="test-operator",
    ))

    decisions = {item.ruleId: item.manualDecision for item in updated.results}
    assert decisions["CASE-001"].status == ManualStatus.RESOLVED
    assert decisions["CASE-002"].status == ManualStatus.PENDING
    assert decisions["CONTACT-001"].status == ManualStatus.IGNORED
    assert decisions["CONTACT-001"].reason == "已人工核对联系方式。"
    events = store.STORE.get_audit_snapshot(task.id)["events"]
    invalidated = next(event for event in events if event["event_type"] == "decisions_invalidated")
    assert invalidated["payload"]["ruleIds"] == ["CASE-002"]


def test_follow_up_cannot_overwrite_archive_completed_while_model_was_running(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "STORE", SqliteTaskStore(tmp_path / "reviews.db"))
    task = _task()
    task.extractionPayload = _complete_extraction().model_dump(mode="json")
    store.save_task(task)
    task.documentId = store.STORE.save_document(
        task.id,
        filename=task.document.name,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        sha256="a" * 64,
        original_path="test/original.docx",
    )
    task.documentVersionId = store.STORE.save_document_version(
        task.documentId,
        content_sha256="a" * 64,
        parsed_payload=task.document.model_dump(mode="json"),
    )
    store.save_task(task)

    async def archive_during_review(*_args, **_kwargs):
        current = store.get_task(task.id)
        current.reviewStatus = ReviewStatus.ARCHIVED
        store.save_task(current)
        return TemplateReviewOutcome(
            extraction=_complete_extraction(),
            issues=[],
            results=task.results,
        )

    monkeypatch.setattr(review, "run_template_review", archive_during_review)

    with pytest.raises(AppError) as error:
        asyncio.run(review.record_follow_up_answer(
            task.id,
            "CASE-001",
            question="你因何事报案？",
            answer="因测试诈骗损失报案。",
            actor_id="test-operator",
        ))

    assert error.value.code == "task_revision_conflict"
    assert store.get_task(task.id).reviewStatus == ReviewStatus.ARCHIVED


def test_failed_domain_retry_restores_completed_state(tmp_path, monkeypatch):
    database = tmp_path / "reviews.db"
    monkeypatch.setattr(store, "STORE", SqliteTaskStore(database))
    task = _task()
    task.status = TaskStatus.FAILED
    task.failedDomains = ["case_timeline"]
    task.extractionPayload = CaseExtraction().model_dump(mode="json")
    store.save_task(task)
    complete = _complete_extraction()
    outcome = TemplateReviewOutcome(
        extraction=complete,
        issues=[],
        results=[
            task.results[0].model_copy(update={"ruleId": f"RULE-{index:03d}"})
            for index in range(34)
        ],
    )
    run = AsyncMock(return_value=outcome)
    monkeypatch.setattr(review, "run_template_review", run)

    updated = asyncio.run(review.retry_failed_domain(task.id, "case_timeline", "test-operator"))

    assert updated.status == TaskStatus.COMPLETED
    assert updated.failedDomains == []
    assert run.await_args.kwargs["domains"] == DOMAIN_ORDER
    assert len(updated.extractionPayload["facts"]) == 92
    assert len(updated.results) == 34
    assert store.STORE.get_audit_snapshot(task.id)["events"][-1]["event_type"] == "domain_retried"


def test_failed_domain_retry_stays_failed_when_required_coverage_is_incomplete(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "STORE", SqliteTaskStore(tmp_path / "reviews.db"))
    task = _task()
    task.status = TaskStatus.FAILED
    task.failedDomains = ["case_timeline"]
    task.extractionPayload = CaseExtraction().model_dump(mode="json")
    store.save_task(task)
    incomplete = _complete_extraction()
    incomplete.facts.pop(DOMAIN_FACT_PATHS["case_timeline"][0])
    monkeypatch.setattr(review, "run_template_review", AsyncMock(return_value=TemplateReviewOutcome(
        extraction=incomplete,
        issues=[],
        results=[],
    )))

    updated = asyncio.run(review.retry_failed_domain(task.id, "case_timeline", "test-operator"))

    assert updated.status == TaskStatus.FAILED
    assert updated.failedDomains == ["case_timeline"]
    assert updated.results == []


def test_upload_persists_original_version_run_fact_issue_and_artifact(tmp_path, monkeypatch):
    database = tmp_path / "reviews.db"
    monkeypatch.setattr(store, "STORE", SqliteTaskStore(database))
    monkeypatch.setattr(artifacts, "ARTIFACT_STORAGE", ArtifactStorage(tmp_path / "artifacts"))
    task = review.create_task(ReviewMode.QWEN)
    document = Document()
    document.add_paragraph("问：你因何事报案？")
    document.add_paragraph("答：因脱敏测试诈骗报案。")
    output = BytesIO()
    document.save(output)
    result = _task().results[0]
    extraction = CaseExtraction(facts={
        "case.report_reason": ExtractedFact(value="脱敏测试诈骗", clarity="clear", evidenceAnchorIds=[]),
    })
    issue = TemplateReviewIssue(
        ruleId="CASE-001",
        ruleName="测试规则",
        group="CASE",
        status=RuleStatus.COVERED,
        reason="已覆盖。",
        suggestedQuestion="",
        severity="high",
    )
    monkeypatch.setattr(review, "run_template_review", AsyncMock(return_value=TemplateReviewOutcome(
        extraction=extraction,
        issues=[issue],
        results=[result],
    )))

    asyncio.run(review.process_upload(task.id, "脱敏测试.docx", output.getvalue()))

    saved = store.get_task(task.id)
    snapshot = store.STORE.get_audit_snapshot(task.id)
    assert saved.status == TaskStatus.COMPLETED
    assert saved.documentId and saved.documentVersionId and saved.reviewRunId
    assert saved.requiredArtifacts == REQUIRED_ARTIFACTS
    assert {artifact.type for artifact in saved.artifacts} == {"original"}
    assert snapshot["document"]["filename"] == "脱敏测试.docx"
    assert len(snapshot["versions"]) == 1
    assert len(snapshot["runs"]) == 1
    assert snapshot["facts"][0]["path"] == "case.report_reason"
    assert snapshot["issues"][0]["rule_id"] == "CASE-001"
    assert snapshot["artifacts"][0]["artifact_type"] == "original"


def test_upload_persists_successful_domains_when_one_domain_fails(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "STORE", SqliteTaskStore(tmp_path / "reviews.db"))
    monkeypatch.setattr(artifacts, "ARTIFACT_STORAGE", ArtifactStorage(tmp_path / "artifacts"))
    task = review.create_task(ReviewMode.QWEN)
    document = Document()
    document.add_paragraph("问：你因何事报案？")
    document.add_paragraph("答：因脱敏测试诈骗报案。")
    output = BytesIO()
    document.save(output)
    partial = _complete_extraction()
    for path in DOMAIN_FACT_PATHS["contact_channels"]:
        partial.facts.pop(path)
    for entity_type in DOMAIN_ENTITY_FIELDS["contact_channels"]:
        partial.entities.pop(entity_type)
    monkeypatch.setattr(review, "run_template_review", AsyncMock(side_effect=extraction.TemplateDomainFailure(
        partial_extraction=partial,
        failed_domains=["contact_channels"],
    )))

    asyncio.run(review.process_upload(task.id, "脱敏测试.docx", output.getvalue()))

    saved = store.get_task(task.id)
    assert saved.status == TaskStatus.FAILED
    assert saved.failedDomains == ["contact_channels"]
    assert set(saved.extractionPayload["facts"]) == set(partial.facts)


def test_openapi_exposes_complete_review_lifecycle_routes():
    paths = app.openapi()["paths"]

    assert {
        "/api/v1/reviews/{task_id}/versions",
        "/api/v1/reviews/{task_id}/issues/{rule_id}/actions",
        "/api/v1/reviews/{task_id}/issues/{rule_id}/follow-up-answer",
        "/api/v1/reviews/{task_id}/domains/{domain}/retry",
        "/api/v1/reviews/{task_id}/warnings/acknowledge",
        "/api/v1/reviews/{task_id}/artifacts/{artifact_id}",
        "/api/v1/reviews/{task_id}/archive",
    } <= set(paths)
