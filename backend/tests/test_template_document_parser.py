import asyncio
import base64
from io import BytesIO
import os
from pathlib import Path
import struct
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
import fitz
from docx import Document
from docx.shared import Inches
from lxml import etree

from app.parsing.parser import parse_document
from app.parsing.openxml import read_docx_parts


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl9sAAAAASUVORK5CYII="
)


def _docx_with_corrupt_media() -> bytes:
    document = Document()
    document.add_paragraph("电信网络诈骗案件询问笔录（模板）")
    document.add_paragraph("问：你听清楚了没有？")
    document.add_paragraph("答：我听清楚了。")
    document.add_picture(BytesIO(PNG_1X1), width=Inches(0.2))
    output = BytesIO()
    document.save(output)

    content = bytearray(output.getvalue())
    with ZipFile(BytesIO(content)) as archive:
        image = archive.getinfo("word/media/image1.png")

    filename_length, extra_length = struct.unpack_from("<HH", content, image.header_offset + 26)
    compressed_start = image.header_offset + 30 + filename_length + extra_length
    content[compressed_start + image.compress_size // 2] ^= 0x01
    return bytes(content)


def _simple_docx(*paragraphs: str) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _docx_with_header_and_footer() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "内部询问笔录页眉"
    document.add_paragraph("正文问答")
    document.sections[0].footer.paragraphs[0].text = "被询问人签名：测试签名"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _rewrite_docx(content: bytes, replacements: dict[str, bytes], additions: dict[str, bytes] | None = None) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item.filename, replacements.get(item.filename, source.read(item.filename)))
        for name, value in (additions or {}).items():
            target.writestr(name, value)
    return output.getvalue()


def _docx_with_body_content_control() -> bytes:
    content = _simple_docx("内容控件中的关键询问事实")
    with ZipFile(BytesIO(content)) as archive:
        xml = archive.read("word/document.xml")
    root = etree.fromstring(xml)
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    body = root.find(f"{{{namespace}}}body")
    paragraph = body.find(f"{{{namespace}}}p")
    index = body.index(paragraph)
    body.remove(paragraph)
    control = etree.Element(f"{{{namespace}}}sdt")
    control_content = etree.SubElement(control, f"{{{namespace}}}sdtContent")
    control_content.append(paragraph)
    body.insert(index, control)
    rewritten = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return _rewrite_docx(content, {"word/document.xml": rewritten})


def _docx_with_document_xml(xml: str) -> bytes:
    content = _simple_docx("placeholder")
    return _rewrite_docx(content, {"word/document.xml": xml.encode("utf-8")})


def _pdf_with_native_text() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 100),
        "Question: describe the fraud timeline. Answer: contact, transfer, and report details.",
    )
    content = document.tobytes()
    document.close()
    return content


def test_docx_with_corrupt_nonessential_media_keeps_document_text():
    parsed = asyncio.run(parse_document("template.docx", _docx_with_corrupt_media()))

    assert "电信网络诈骗案件询问笔录" in parsed.text
    assert any(warning.code == "media_corrupt" for warning in parsed.warnings)


def test_document_blocks_have_stable_ids_and_exact_character_ranges():
    content = _simple_docx("第一段模板文本", "第二段模板文本")

    first = asyncio.run(parse_document("template.docx", content))
    second = asyncio.run(parse_document("template.docx", content))

    first_blocks = [paragraph for page in first.pages for paragraph in page.paragraphs]
    second_blocks = [paragraph for page in second.pages for paragraph in page.paragraphs]
    assert [block.id for block in first_blocks] == [block.id for block in second_blocks]
    assert all(block.id for block in first_blocks)
    assert all(first.text[block.charStart:block.charEnd] == block.text for block in first_blocks)


def test_document_block_ids_ignore_nonvisible_docx_metadata():
    def docx_with_title(title: str) -> bytes:
        document = Document()
        document.core_properties.title = title
        document.add_paragraph("相同的可见询问笔录正文")
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    first = asyncio.run(parse_document("record.docx", docx_with_title("元数据版本一")))
    second = asyncio.run(parse_document("record.docx", docx_with_title("元数据版本二")))

    assert first.text == second.text
    assert first.pages[0].paragraphs[0].id == second.pages[0].paragraphs[0].id


def test_docx_parser_preserves_header_and_footer_text():
    parsed = asyncio.run(parse_document("template.docx", _docx_with_header_and_footer()))

    assert "内部询问笔录页眉" in parsed.text
    assert "正文问答" in parsed.text
    assert "被询问人签名：测试签名" in parsed.text


def test_docx_parser_preserves_body_content_controls_in_reading_order():
    parsed = asyncio.run(parse_document("template.docx", _docx_with_body_content_control()))

    assert parsed.text == "内容控件中的关键询问事实"


def test_docx_parser_transcribes_text_checkbox_glyphs_in_question_answer_order():
    content = _simple_docx(
        "问：你通过哪些方式联系？",
        "答：☑ 电话 ☐ 短信 ☒ APP □ 其他",
    )

    parsed = asyncio.run(parse_document("record.docx", content))

    assert parsed.questionAnswers[0].answer == "[选中] 电话 [未选] 短信 [选中] APP [未选] 其他"


def test_docx_parser_preserves_wingdings_checkbox_symbols():
    xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:body>
        <w:p><w:r><w:t>问：你选择哪些渠道？</w:t></w:r></w:p>
        <w:p>
          <w:r><w:t>答：</w:t></w:r>
          <w:r><w:sym w:font="Wingdings 2" w:char="00A3"/></w:r>
          <w:r><w:t> 电话 </w:t></w:r>
          <w:r><w:sym w:font="Wingdings 2" w:char="0052"/></w:r>
          <w:r><w:t> APP</w:t></w:r>
        </w:p>
        <w:sectPr/>
      </w:body>
    </w:document>"""

    parsed = asyncio.run(parse_document("record.docx", _docx_with_document_xml(xml)))

    assert parsed.questionAnswers[0].answer == "[未选] 电话 [选中] APP"


def test_docx_parser_transcribes_content_control_and_legacy_form_checkboxes():
    xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <w:document
        xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">
      <w:body>
        <w:p><w:r><w:t>问：你选择哪些渠道？</w:t></w:r></w:p>
        <w:p>
          <w:r><w:t>答：</w:t></w:r>
          <w:sdt>
            <w:sdtPr><w14:checkbox><w14:checked w14:val="1"/></w14:checkbox></w:sdtPr>
            <w:sdtContent><w:r><w:t>☒</w:t></w:r></w:sdtContent>
          </w:sdt>
          <w:r><w:t> 电话 </w:t></w:r>
          <w:r><w:fldChar w:fldCharType="begin"><w:ffData><w:checkBox><w:default w:val="0"/></w:checkBox></w:ffData></w:fldChar></w:r>
          <w:r><w:t> 短信</w:t></w:r>
        </w:p>
        <w:sectPr/>
      </w:body>
    </w:document>"""

    parsed = asyncio.run(parse_document("record.docx", _docx_with_document_xml(xml)))

    assert parsed.questionAnswers[0].answer == "[选中] 电话 [未选] 短信"


def test_docx_parser_keeps_unknown_checkbox_state_and_emits_warning():
    xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:body>
        <w:p><w:r><w:t>问：你选择哪些渠道？</w:t></w:r></w:p>
        <w:p>
          <w:r><w:t>答：</w:t></w:r>
          <w:r><w:sym w:font="Unknown Checkbox Font" w:char="FFFF"/></w:r>
          <w:r><w:t> 其他</w:t></w:r>
        </w:p>
        <w:sectPr/>
      </w:body>
    </w:document>"""

    parsed = asyncio.run(parse_document("record.docx", _docx_with_document_xml(xml)))

    assert parsed.questionAnswers[0].answer == "[状态不明] 其他"
    assert any(warning.code == "checkbox_state_unknown" for warning in parsed.warnings)


def test_docx_reader_ignores_unreferenced_hidden_header_and_media_parts():
    content = _simple_docx("可见正文")
    hidden_header = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:p><w:r><w:t>隐藏孤立页眉</w:t></w:r></w:p></w:hdr>'
    ).encode("utf-8")
    injected = _rewrite_docx(content, {}, {
        "word/header99.xml": hidden_header,
        "word/media/orphan.png": PNG_1X1,
    })

    parts = read_docx_parts(injected)

    assert [block.text for block in parts.blocks] == ["可见正文"]
    assert parts.images == []


def test_docx_reader_rejects_suspicious_package_expansion():
    content = _simple_docx("可见正文")
    with ZipFile(BytesIO(content)) as archive:
        xml = archive.read("word/document.xml")
    expanded = xml.replace(b"</w:body>", b"A" * (2 * 1024 * 1024) + b"</w:body>")
    suspicious = _rewrite_docx(content, {"word/document.xml": expanded})

    with pytest.raises(Exception) as error:
        read_docx_parts(suspicious)

    assert getattr(error.value, "code", None) == "docx_expansion_limit"


def test_parsed_document_exposes_reconstructed_question_answers():
    content = _simple_docx(
        "问：是否收到风险提示？",
        "答：收到银行短信提示。",
    )

    parsed = asyncio.run(parse_document("record.docx", content))

    assert parsed.text == "问：是否收到风险提示？\n答：收到银行短信提示。"
    assert len(parsed.pages) == 1
    assert [(block.question, block.answer) for block in parsed.questionAnswers] == [
        ("是否收到风险提示？", "收到银行短信提示。"),
    ]
    assert parsed.questionAnswers[0].anchorIds == [
        parsed.pages[0].paragraphs[0].id,
        parsed.pages[0].paragraphs[1].id,
    ]


def test_pdf_native_text_blocks_preserve_source_coordinates():
    parsed = asyncio.run(parse_document("record.pdf", _pdf_with_native_text()))

    block = parsed.pages[0].paragraphs[0]
    assert block.bbox is not None
    x0, y0, x1, y1 = block.bbox
    assert 0 <= x0 < x1
    assert 0 <= y0 < y1


@pytest.mark.skipif(not os.getenv("BILU_TEMPLATE_PATH"), reason="BILU_TEMPLATE_PATH is not configured")
def test_provided_internal_template_is_parseable_when_available():
    path = Path(os.environ["BILU_TEMPLATE_PATH"])

    parsed = asyncio.run(parse_document(path.name, path.read_bytes()))

    assert "电信网络诈骗案件询问笔录" in parsed.text
    body_blocks = [
        paragraph
        for page in parsed.pages
        for paragraph in page.paragraphs
        if paragraph.sourceType.value in {"native_text", "table"}
    ]
    assert len(body_blocks) == 144
    assert parsed.text.count("问：") + parsed.text.count("问:") == 33
    assert len(parsed.questionAnswers) == 33
    basic_information = next(
        block for block in parsed.questionAnswers if "基本情况" in block.question
    )
    assert basic_information.answer == ""
    assert any("事主基本信息" in item for item in basic_information.guidance)
    assert "以上笔录给你看一下" in parsed.text
    assert any(warning.code == "media_corrupt" for warning in parsed.warnings)
