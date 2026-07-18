from pathlib import Path
import re

from docx import Document

from scripts.generate_test_records import generate_test_records


EXPECTED_FILES = [
    "01-基本完整-电诈询问笔录.docx",
    "02-明确漏问-电诈询问笔录.docx",
    "03-复杂场景-APP返利询问笔录.docx",
    "04-冒充公检法-老年人ATM询问笔录.docx",
    "05-投资交友-数字货币询问笔录.docx",
    "06-冒充客服-远程控制询问笔录.docx",
    "07-线下取现-寄递黄金询问笔录.docx",
]

REQUIRED_IDENTITY_LABELS = [
    "姓名",
    "性别",
    "年龄",
    "民族",
    "身份证号码",
    "工作单位",
    "现住址",
    "联系电话",
]


def _document_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_generate_seven_formal_fictional_records(tmp_path):
    generated = generate_test_records(tmp_path)

    assert [path.name for path in generated] == EXPECTED_FILES
    case_numbers = set()
    for path in generated:
        assert path.exists()
        text = _document_text(path)
        assert "询问笔录" in text
        assert "虚构测试材料" in text
        assert "案件编号" in text
        assert "权利义务告知" in text
        assert "被询问人签名" in text
        assert not re.search(r"(?<!\d)1\d{16}[0-9X](?!\d)", text)
        assert not re.search(r"(?<!\d)1[3-9]\d{9}(?!\d)", text)
        assert not re.search(r"(?<!\d)(?:62\d{14,17}|[3-6]\d{15})(?!\d)", text)
        assert not re.search(r"https?://", text, re.IGNORECASE)
        assert not re.search(r"(?<!\d)(?:(?:[1-9]|[1-9]\d|1\d\d|2[0-4]\d|25[0-5])\.){3}(?:[1-9]|[1-9]\d|1\d\d|2[0-4]\d|25[0-5])(?!\d)", text)
        assert re.search(r"(?:流水号|交易单号)[:：][A-Z0-9-]{8,}", text)
        assert not any(token in text for token in ["待补充", "XXX", "____", "不详"])
        assert all(label in text for label in REQUIRED_IDENTITY_LABELS)
        case_numbers.update(re.findall(r"TEST-2026-0716-\d{3}", text))
    assert len(case_numbers) == 7


def test_unknown_information_is_an_explicit_answer_only(tmp_path):
    generated = generate_test_records(tmp_path)
    texts = {path.name: _document_text(path) for path in generated}

    assert "我不知道" not in texts[EXPECTED_FILES[0]]
    assert "我不知道" not in texts[EXPECTED_FILES[1]]
    assert "答：我不知道" in texts[EXPECTED_FILES[2]]
    for text in texts.values():
        for line in text.splitlines():
            if "不知道" in line:
                assert "答：我不知道" in line


def test_answers_do_not_repeat_the_answer_prefix(tmp_path):
    generated = generate_test_records(tmp_path)

    for path in generated:
        assert "答：答：" not in _document_text(path), path.name


def test_each_record_contains_scenario_specific_safe_namespaces(tmp_path):
    generated = generate_test_records(tmp_path)
    texts = {path.name: _document_text(path) for path in generated}

    assert "ACCOUNT_WX_CASE_01" in texts[EXPECTED_FILES[0]]
    assert "ACCOUNT_QQ_CASE_02" in texts[EXPECTED_FILES[1]]
    assert "ACCOUNT_APP_CASE_03" in texts[EXPECTED_FILES[2]]
    assert "IP_TEST_CASE_03" in texts[EXPECTED_FILES[2]]
    assert "URL_TEST_CASE_03" in texts[EXPECTED_FILES[2]]


def test_corpus_covers_high_value_telecom_fraud_edges(tmp_path):
    generated = generate_test_records(tmp_path)
    texts = {path.name: _document_text(path) for path in generated}

    assert all(token in texts[EXPECTED_FILES[3]] for token in ["安全账户", "ATM", "通话记录", "删除"])
    assert all(token in texts[EXPECTED_FILES[4]] for token in ["投资", "数字货币", "USDT", "前后陈述"])
    assert all(token in texts[EXPECTED_FILES[5]] for token in ["屏幕共享", "远程控制", "验证码", "网贷"])
    assert all(token in texts[EXPECTED_FILES[6]] for token in ["取现", "寄递黄金", "网约车", "未询问"])
