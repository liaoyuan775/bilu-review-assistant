import asyncio
import base64
from io import BytesIO
from unittest.mock import AsyncMock, patch

import fitz
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
import pytest

from app.core.errors import AppError
from app.main import app
from app.data.rules import TEMPLATE_RULES
from app.data.demo_cases import DEMO_CASES, load_demo_document
from app.storage import artifacts
from app.storage.artifacts import ArtifactStorage
from app.review.extraction import TemplateReviewOutcome
from app.review.mock import build_mock_review
from app.core.template_models import CaseExtraction
from app.storage import store
from app.storage.store import SqliteTaskStore


client = TestClient(app)


def _demo_document(index: int = 1):
    return asyncio.run(load_demo_document(DEMO_CASES[index]))


@pytest.fixture(autouse=True)
def isolated_review_store(tmp_path, monkeypatch):
    monkeypatch.setattr(store, "STORE", SqliteTaskStore(tmp_path / "reviews.db"))
    monkeypatch.setattr(artifacts, "ARTIFACT_STORAGE", ArtifactStorage(tmp_path / "artifacts"))


def _docx_bytes(paragraphs: list[str]) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_with_table_bytes() -> bytes:
    document = Document()
    document.add_paragraph("问：请说明转账情况。答：我通过手机银行转账。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "转账时间"
    table.cell(0, 1).text = "转账金额"
    table.cell(1, 0).text = "2026年7月15日10时30分"
    table.cell(1, 1).text = "5000元"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_with_image_bytes() -> bytes:
    image = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    document = Document()
    document.add_paragraph("问：是否保留转账凭证？答：凭证截图附后。")
    document.add_picture(BytesIO(image))
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _text_pdf_bytes(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(payload)


def _blank_pdf_bytes(page_count: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=612, height=792)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _mixed_pdf_bytes() -> bytes:
    image = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Native interview text retained for evidence location.")
    page.insert_image(fitz.Rect(72, 100, 180, 208), stream=image)
    content = document.tobytes()
    document.close()
    return content


def _upload(filename: str, content: bytes):
    async def fixture_review(document, group_timings=None):
        if group_timings is not None:
            group_timings.update({"三现": 12, "四流": 18})
        return TemplateReviewOutcome(
            extraction=CaseExtraction(),
            issues=[],
            results=build_mock_review(document),
        )

    with patch("app.review.review.run_template_review", side_effect=fixture_review):
        created = client.post(
            "/api/v1/reviews",
            data={"mode": "local"},
            files={"file": (filename, content, "application/octet-stream")},
        )
    assert created.status_code == 202
    task = client.get(f"/api/v1/reviews/{created.json()['taskId']}").json()
    assert task["mode"] == "qwen"
    return task


def test_upload_records_stage_and_group_timings():
    task = _upload("timed.docx", _docx_bytes(["问：案发经过？答：脱敏测试内容。"]))

    assert task["timings"]["parseMs"] >= 0
    assert task["timings"]["modelReviewMs"] >= 0
    assert task["timings"]["modelGroupsMs"] == {"三现": 12, "四流": 18}
    assert task["timings"]["totalMs"] >= task["timings"]["modelReviewMs"]


def test_health_and_rules():
    health = client.get("/api/v1/health")
    assert health.status_code == 200
    assert health.json()["ruleCount"] == len(TEMPLATE_RULES)
    rules = client.get("/api/v1/rules")
    assert rules.status_code == 200
    payload = rules.json()["rules"]
    assert len(payload) == len(TEMPLATE_RULES)
    assert {rule["group"] for rule in payload} == {rule.group for rule in TEMPLATE_RULES}
    assert all(rule["source"].startswith("内部询问笔录模板 v") for rule in payload)
    assert all(rule["requiredFacts"] for rule in payload)

    schema = client.get("/api/openapi.json").json()
    task_statuses = schema["components"]["schemas"]["TaskStatus"]["enum"]
    assert task_statuses == ["uploading", "parsing", "recognizing", "checking", "validating", "completed", "failed"]
    review_response = schema["paths"]["/api/v1/reviews/{task_id}"]["get"]["responses"]["200"]
    assert review_response["content"]["application/json"]["schema"]["$ref"].endswith("/ReviewTask")


def test_demo_review_decisions_block_archive_until_artifacts_exist():
    created = client.post("/api/v1/reviews/demos/case-01-baseline", json={})
    assert created.status_code == 202
    task_id = created.json()["taskId"]
    task = client.get(f"/api/v1/reviews/{task_id}").json()
    assert task["status"] == "completed"
    assert len(task["results"]) == len(TEMPLATE_RULES)
    assert task["requiredArtifacts"] == ["review_pdf", "follow_up_docx", "structured_json", "archive_manifest"]
    assert [artifact["type"] for artifact in task["artifacts"]] == ["original"]
    blocked = client.post(f"/api/v1/reviews/{task_id}/complete")
    assert blocked.status_code == 409

    actionable = [item for item in task["results"] if item["status"] in {"missing", "incomplete", "inconsistent", "needs_manual_review"}]
    for item in actionable:
        saved = client.post(
            f"/api/v1/reviews/{task_id}/issues/{item['ruleId']}/actions",
            json={"status": "resolved", "reason": "已通过脱敏测试补问核对。", "actorId": "test-operator"},
        )
        assert saved.status_code == 200

    still_blocked = client.post(f"/api/v1/reviews/{task_id}/complete")
    assert still_blocked.status_code == 409
    assert still_blocked.json()["error"]["code"] == "archive_missing_artifacts"

    follow_ups = client.get(f"/api/v1/reviews/{task_id}/follow-ups")
    assert follow_ups.status_code == 200
    assert follow_ups.json()["items"] == []

    report = client.get(f"/api/v1/reviews/{task_id}/report-data")
    assert report.status_code == 200
    assert report.json()["reviewStatus"] == "in_review"
    assert report.json()["victimProfile"] == task["victimProfile"]
    assert len(report.json()["results"]) == len(TEMPLATE_RULES)

    history = client.get("/api/v1/reviews")
    assert history.status_code == 200
    assert any(item["id"] == task_id and item["reviewStatus"] == "in_review" for item in history.json()["reviews"])


def test_demo_task_records_demo_id():
    created = client.post("/api/v1/reviews/demos/case-01-baseline", json={})
    task = client.get(f"/api/v1/reviews/{created.json()['taskId']}").json()

    assert task["demoId"] == "case-01-baseline"


def test_bulk_demo_pass_generates_artifacts_and_leaves_archive_ready():
    created = client.post("/api/v1/reviews/demos/case-01-baseline", json={}).json()
    response = client.post(f"/api/v1/reviews/{created['taskId']}/demo-pass")

    assert response.status_code == 200
    task = response.json()
    actionable = {"missing", "incomplete", "inconsistent", "needs_manual_review"}
    handled = [item for item in task["results"] if item["status"] in actionable]
    assert handled
    assert all(item["manualDecision"]["status"] == "resolved" for item in handled)
    assert all(item["manualDecision"]["reason"] == "脱敏演示：一键测试通过" for item in handled)
    assert set(task["acknowledgedWarnings"]) == {
        warning["code"] for warning in task["document"]["warnings"]
    }
    assert {artifact["type"] for artifact in task["artifacts"]} >= {
        "review_pdf", "follow_up_docx", "structured_json", "archive_manifest",
    }
    archived = client.post(f"/api/v1/reviews/{created['taskId']}/archive")
    assert archived.status_code == 200


def test_bulk_demo_pass_rejects_normal_upload():
    task = _upload("normal.docx", _docx_bytes(["问：案发经过？答：脱敏测试内容。"]))

    response = client.post(f"/api/v1/reviews/{task['id']}/demo-pass")

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "demo_action_not_allowed"


def test_bulk_demo_pass_rejects_failed_domains():
    created = client.post("/api/v1/reviews/demos/case-01-baseline", json={}).json()
    task = store.get_task(created["taskId"])
    task.failedDomains = ["content"]
    store.save_task(task)

    response = client.post(f"/api/v1/reviews/{created['taskId']}/demo-pass")

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "demo_has_failed_domains"


def test_template_results_locate_source_paragraphs():
    created = client.post("/api/v1/reviews/demos/case-01-baseline", json={})
    assert created.status_code == 202
    task = client.get(f"/api/v1/reviews/{created.json()['taskId']}").json()
    for item in task["results"]:
        location = item["evidenceLocation"]
        assert location is not None
        page = task["document"]["pages"][location["page"] - 1]
        assert 1 <= location["paragraph"] <= len(page["paragraphs"])


def test_docx_upload_completes_with_all_four_state_results():
    task = _upload("脱敏笔录.docx", _docx_bytes([
        "问：请说明案发时间、地点和经过。答：2026年7月15日，我在测试路附近接到陌生电话。",
        "问：对方使用什么方式联系？答：通过电话联系，让我点击一个链接并转账。",
        "问：损失金额和支付方式？答：通过银行卡转账5000元，交易流水号为TEST-001。",
        "问：是否保存聊天记录和转账凭证？答：相关截图已经保存。",
    ]))
    assert task["status"] == "completed"
    assert task["document"]["format"] == "DOCX"
    assert len(task["results"]) == len(TEMPLATE_RULES)
    assert {item["status"] for item in task["results"]} <= {"covered", "missing", "incomplete", "not_applicable"}


def test_text_pdf_upload_completes():
    text = (
        "Interview record for anonymized demo. The incident occurred on July 15 at Test Road. "
        "The caller requested a bank transfer. The transaction reference is TEST-002 and screenshots were retained."
    )
    task = _upload("anonymized-record.pdf", _text_pdf_bytes(text))
    assert task["status"] == "completed"
    assert task["document"]["format"] == "PDF"
    assert len(task["results"]) == len(TEMPLATE_RULES)


def test_scanned_pdf_uses_multimodal_transcription(monkeypatch):
    transcribe = AsyncMock(return_value={
        "paragraphs": ["问：转账金额是多少？答：5000元。"],
        "confidence": 0.96,
    })
    monkeypatch.setattr("app.parsing.parser.transcribe_image", transcribe, raising=False)

    task = _upload("scanned.pdf", _blank_pdf_bytes())

    assert task["status"] == "completed"
    assert transcribe.await_count == 1
    paragraph = task["document"]["pages"][0]["paragraphs"][0]
    assert paragraph["text"] == "问：转账金额是多少？答：5000元。"
    assert paragraph["sourceType"] == "vision"
    assert paragraph["confidence"] == 0.96
    assert paragraph["id"]
    assert paragraph["charStart"] == 0
    assert paragraph["charEnd"] == len(paragraph["text"])
    assert paragraph["bbox"] is None


def test_mixed_pdf_keeps_native_text_and_transcribes_embedded_images(monkeypatch):
    transcribe = AsyncMock(return_value={
        "paragraphs": ["图片凭证流水号：MIXED-001"],
        "confidence": 0.9,
    })
    monkeypatch.setattr("app.parsing.parser.transcribe_image", transcribe)

    task = _upload("mixed-record.pdf", _mixed_pdf_bytes())

    assert task["status"] == "completed"
    blocks = task["document"]["pages"][0]["paragraphs"]
    assert {block["sourceType"] for block in blocks} == {"native_text", "vision"}
    assert any("MIXED-001" in block["text"] for block in blocks)


def test_docx_tables_are_preserved_as_structured_blocks():
    task = _upload("table-record.docx", _docx_with_table_bytes())

    assert task["status"] == "completed"
    blocks = [paragraph for page in task["document"]["pages"] for paragraph in page["paragraphs"]]
    assert any(block["sourceType"] == "table" and "2026年7月15日10时30分" in block["text"] for block in blocks)


def test_docx_embedded_images_use_multimodal_transcription(monkeypatch):
    transcribe = AsyncMock(return_value={
        "paragraphs": ["交易流水号：TEST-VISION-001"],
        "confidence": 0.91,
    })
    monkeypatch.setattr("app.parsing.parser.transcribe_image", transcribe, raising=False)

    task = _upload("image-record.docx", _docx_with_image_bytes())

    assert task["status"] == "completed"
    assert transcribe.await_count == 1
    blocks = [paragraph for page in task["document"]["pages"] for paragraph in page["paragraphs"]]
    assert any(block["sourceType"] == "vision" and "TEST-VISION-001" in block["text"] for block in blocks)


def test_legacy_doc_is_rejected_with_conversion_guidance():
    task = _upload("legacy-record.doc", b"legacy-word-content")

    assert task["status"] == "failed"
    assert task["errorCode"] == "unsupported_legacy_word"
    assert "DOCX" in task["errorMessage"]


def test_long_docx_is_accepted_without_page_limit():
    paragraphs = [f"第 {index + 1} 部分。" + "脱敏测试内容" * 300 for index in range(11)]
    task = _upload("long-record.docx", _docx_bytes(paragraphs))
    assert task["status"] == "completed"
    assert task["document"]["pageCount"] >= 11
    assert len(task["results"]) == len(TEMPLATE_RULES)


def test_empty_and_corrupt_files_fail_cleanly():
    empty = _upload("empty.docx", b"")
    assert empty["status"] == "failed"
    assert empty["errorCode"] == "empty_document"
    assert empty["results"] == []

    corrupt = _upload("corrupt.pdf", b"not a pdf")
    assert corrupt["status"] == "failed"
    assert corrupt["errorCode"] == "parse_failed"
    assert corrupt["results"] == []


@pytest.mark.parametrize(
    ("failure", "expected_code"),
    [
        (AppError("model_unreachable", "模型服务不可达。", 503), "model_unreachable"),
        (RuntimeError("unexpected model adapter failure"), "review_service_failed"),
    ],
)
def test_qwen_failures_fail_the_whole_demo_task(monkeypatch, failure, expected_code):
    monkeypatch.setattr("app.review.review.run_template_review", AsyncMock(side_effect=failure))
    created = client.post("/api/v1/reviews/demos/case-02-line-breaks", json={})
    assert created.status_code == 202
    task = client.get(f"/api/v1/reviews/{created.json()['taskId']}").json()
    assert task["status"] == "failed"
    assert task["errorCode"] == expected_code
    assert task["results"] == []
